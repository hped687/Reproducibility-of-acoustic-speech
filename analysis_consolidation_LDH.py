#!/usr/bin/env python3
"""Create analysis-ready datasets, descriptive tables, figures, and an Excel summary
for the systematic review of speech-feature extraction in psychiatry.


Example
-------
python systematic_review_analysis_pipeline.py \
  --studies "review_719035_20260724155533.csv" \
  --github-workbook "github_repository_review_workbook_quality_reviewed.xlsx" \
  --repo-overrides "repository_analysis_overrides.csv" \
  --output-dir "systematic_review_analysis_outputs"
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# -----------------------------------------------------------------------------
# Configuration
# -----------------------------------------------------------------------------

MISSING_MARKERS = {
    "", "nan", "none", "null", "n/a", "na", "not applicable", "not available",
    "missing", "unknown",
}

NOT_REPORTED_PATTERNS = (
    r"\bnot reported\b",
    r"\bnot specified\b",
    r"\bnot stated\b",
    r"\bnot described\b",
    r"\bnot provided\b",
    r"\bunclear\b",
    r"\bunknown\b",
    r"\binsufficient detail",
    r"\bno information",
)

# Ordered dictionaries make the output stable and interpretable.
DIAGNOSIS_PATTERNS: OrderedDict[str, Sequence[str]] = OrderedDict([
    ("Depressive disorders", [r"\bmdd\b", r"major depressive", r"\bdepress(?:ion|ive|ed)?\b", r"dysthymi"]),
    ("Bipolar disorders", [r"\bbipolar\b", r"\bmania\b", r"\bmanic\b"]),
    ("Schizophrenia spectrum/psychosis", [
        r"schizophren", r"schizoaffective", r"\bpsychosis\b", r"\bpsychotic\b",
        r"first[- ]episode psych", r"clinical high risk", r"\bchr[- ]?p\b", r"ultra[- ]high risk", r"\buhr\b",
    ]),
    ("Anxiety disorders/symptoms", [r"\banxiety\b", r"\bgad\b", r"panic disorder", r"social anxiety", r"phobi"]),
    ("OCD and related disorders", [r"\bocd\b", r"obsessive[- ]compulsive", r"body dysmorphic", r"\bbdd\b", r"trichotilloman", r"hoarding"]),
    ("Trauma- and stressor-related disorders", [r"\bptsd\b", r"post[- ]?traumatic", r"trauma exposure", r"childhood trauma"]),
    ("Autism spectrum", [r"\bautism\b", r"\bautistic\b", r"\basd\b"]),
    ("ADHD", [r"\badhd\b", r"attention[- ]deficit"]),
    ("Eating disorders", [r"\banorexi", r"\bbulimi", r"eating disorder", r"binge[- ]?eating"]),
    ("Personality disorders", [r"personality disorder", r"borderline personality"]),
    ("Other/mixed psychiatric", [r"mood disorder", r"psychological distress", r"serious mental illness", r"transdiagnostic", r"mixed psychiatric"]),
])

SOFTWARE_PATTERNS: OrderedDict[str, dict[str, Any]] = OrderedDict([
    ("openSMILE", {"patterns": [r"opensmile", r"open smile"], "type": "Feature-extraction tool"}),
    ("Praat/Parselmouth", {"patterns": [r"\bpraat\b", r"parselmouth"], "type": "Feature-extraction tool"}),
    ("librosa", {"patterns": [r"\blibrosa\b"], "type": "Feature-extraction tool"}),
    ("MATLAB", {"patterns": [r"\bmatlab\b"], "type": "Programming/modelling environment"}),
    ("COVAREP", {"patterns": [r"\bcovarep\b"], "type": "Feature-extraction tool"}),
    ("Kaldi", {"patterns": [r"\bkaldi\b"], "type": "Speech-processing toolkit"}),
    ("pyAudioAnalysis", {"patterns": [r"pyaudioanalysis"], "type": "Feature-extraction tool"}),
    ("VoiceSauce", {"patterns": [r"voicesauce"], "type": "Feature-extraction tool"}),
    ("DisVoice", {"patterns": [r"disvoice"], "type": "Feature-extraction tool"}),
    ("Essentia", {"patterns": [r"\bessentia\b"], "type": "Feature-extraction tool"}),
    ("OpenDBM", {"patterns": [r"opendbm"], "type": "Feature-extraction tool"}),
    ("OpenEAR", {"patterns": [r"openear"], "type": "Feature-extraction tool"}),
    ("SpeechBrain", {"patterns": [r"speechbrain"], "type": "Speech-processing toolkit"}),
    ("TorchAudio", {"patterns": [r"torchaudio"], "type": "Speech-processing toolkit"}),
    ("PyTorch", {"patterns": [r"\bpytorch\b", r"torch\b"], "type": "Programming/modelling framework"}),
    ("TensorFlow/Keras", {"patterns": [r"tensorflow", r"\bkeras\b"], "type": "Programming/modelling framework"}),
    ("scikit-learn", {"patterns": [r"scikit[- ]learn", r"sklearn"], "type": "Programming/modelling framework"}),
    ("R packages", {"patterns": [r"\br package", r"\brstudio\b", r"\br software\b"], "type": "Programming/modelling environment"}),
    ("Wav2Vec/HuBERT/WavLM", {"patterns": [r"wav2vec", r"hubert", r"wavlm"], "type": "Learned speech representation"}),
    ("Custom/in-house/proprietary", {"patterns": [r"custom", r"in[- ]house", r"proprietary", r"internally developed", r"bespoke"], "type": "Custom/proprietary"}),
])

PROGRAMMING_LANGUAGE_PATTERNS: OrderedDict[str, Sequence[str]] = OrderedDict([
    ("Python", [r"\bpython\b", r"jupyter"]),
    ("R", [r"(^|[;,/ ]+)r($|[;,/ .]+)", r"\brstudio\b"]),
    ("MATLAB", [r"\bmatlab\b"]),
    ("C/C++", [r"\bc\+\+\b", r"\bcpp\b", r"\bc language\b"]),
    ("Java", [r"\bjava\b"]),
    ("JavaScript/TypeScript", [r"javascript", r"typescript", r"node\.js"]),
    ("Julia", [r"\bjulia\b"]),
    ("Shell", [r"\bbash\b", r"shell script"]),
])

FEATURE_PATTERNS: OrderedDict[str, dict[str, Any]] = OrderedDict([
    ("Prosody/F0/pitch/rhythm", {"patterns": [r"\bprosod", r"fundamental frequency", r"\bf0\b", r"\bpitch\b", r"intonation", r"rhythm"], "scope": "Acoustic"}),
    ("Temporal/pausing/turn-taking", {"patterns": [r"\bpause", r"speech rate", r"speaking rate", r"articulation rate", r"duration", r"temporal", r"turn[- ]taking", r"response latency", r"voice activity", r"voiced proportion"], "scope": "Acoustic"}),
    ("Energy/intensity/loudness", {"patterns": [r"\benergy\b", r"\bintensity\b", r"\bloudness\b", r"amplitude", r"rms"], "scope": "Acoustic"}),
    ("Cepstral/MFCC/filterbank", {"patterns": [r"\bmfcc", r"cepstr", r"filterbank", r"mel[- ]frequency"], "scope": "Acoustic"}),
    ("Spectral", {"patterns": [r"\bspectral", r"spectrum", r"spectral centroid", r"spectral flux", r"spectral roll", r"spectral slope"], "scope": "Acoustic"}),
    ("Voice quality/perturbation", {"patterns": [r"\bjitter\b", r"\bshimmer\b", r"harmonic[- ]to[- ]noise", r"\bhnr\b", r"cepstral peak prominence", r"\bcpp\b", r"voice quality", r"perturbation", r"harmonicity"], "scope": "Acoustic"}),
    ("Formant/articulatory", {"patterns": [r"\bformant", r"\bf1\b", r"\bf2\b", r"\bf3\b", r"vowel space", r"articulat", r"vocal tract"], "scope": "Acoustic"}),
    ("Glottal/phonation", {"patterns": [r"glottal", r"phonation", r"open quotient", r"closing quotient", r"normalized amplitude quotient", r"naq\b"], "scope": "Acoustic"}),
    ("Spectrogram/time-frequency", {"patterns": [r"spectrogram", r"time[- ]frequency", r"mel spectrogram", r"stft"], "scope": "Acoustic representation"}),
    ("Learned/self-supervised representations", {"patterns": [r"wav2vec", r"hubert", r"wavlm", r"self[- ]supervised", r"learned representation", r"deep embedding", r"speech embedding", r"latent representation", r"transformer embedding"], "scope": "Learned representation"}),
    ("Nonlinear/complexity", {"patterns": [r"entropy", r"fractal", r"nonlinear", r"lyapunov", r"recurrence", r"complexity"], "scope": "Acoustic"}),
    ("Linguistic/text", {"patterns": [r"linguistic", r"semantic", r"syntactic", r"lexical", r"word embedding", r"transcript", r"text feature", r"bert"], "scope": "Non-acoustic/multimodal"}),
])

FEATURE_SET_PATTERNS: OrderedDict[str, Sequence[str]] = OrderedDict([
    ("eGeMAPS/GeMAPS", [r"egemaps", r"\bgemaps\b"]),
    ("ComParE", [r"\bcompare\b", r"interspeech 2013 computational paralinguistics"]),
    ("INTERSPEECH 2009", [r"interspeech 2009", r"is09"]),
    ("INTERSPEECH 2010", [r"interspeech 2010", r"is10"]),
    ("EmoBase", [r"emobase"]),
])

TASK_PATTERNS: OrderedDict[str, Sequence[str]] = OrderedDict([
    ("Clinical interview/conversation", [r"clinical interview", r"semi[- ]structured interview", r"structured interview", r"conversation", r"dialogue", r"discussion", r"telephone call", r"phone call"]),
    ("Read/fixed text", [r"read passage", r"reading task", r"read speech", r"fixed text", r"northwind", r"rainbow passage", r"sentence reading", r"word reading", r"read aloud"]),
    ("Spontaneous/free/narrative speech", [r"free speech", r"spontaneous", r"narrative", r"open[- ]ended", r"freeform", r"picture description", r"story", r"monologue", r"discourse"]),
    ("Sustained vowel/phonation", [r"vowel sound", r"sustained vowel", r"phonation", r"sustain(?:ed)? /?[aeiou]"]),
    ("Vocal diary/EMA/remote sampling", [r"vocal diar", r"voice diar", r"ema", r"ecological momentary", r"smartphone", r"remote", r"interactive voice response", r"daily voice"]),
    ("Verbal fluency/naming/cognitive task", [r"verbal fluency", r"picture naming", r"stroop", r"colour naming", r"color naming", r"cognitive task", r"category fluency", r"letter fluency"]),
    ("Affective/emotion-elicitation task", [r"emotional word", r"affective", r"emotion elic", r"emotion production", r"positive.*negative.*neutral"]),
])

VALIDATION_PATTERNS: OrderedDict[str, Sequence[str]] = OrderedDict([
    ("Cross-validation", [r"cross[- ]validation", r"\bk[- ]?fold\b", r"leave[- ]one[- ]out", r"\bloocv\b", r"nested cv", r"repeated cv"]),
    ("Held-out/train-test evaluation", [r"held[- ]out", r"train(?:ing)?[- /]test", r"test set", r"validation set", r"independent test split"]),
    ("Participant-independent evaluation", [r"subject[- ]independent", r"participant[- ]independent", r"speaker[- ]independent", r"grouped split", r"leave[- ]one[- ]subject"]),
    ("External/independent cohort", [r"external validation", r"external cohort", r"independent cohort", r"independent dataset", r"generalization", r"generalisation", r"cross[- ]dataset", r"replication cohort"]),
    ("Test-retest/reliability", [r"test[- ]retest", r"reliability", r"inter[- ]rater", r"intra[- ]class correlation", r"\bicc\b"]),
    ("Replication", [r"\breplication\b", r"replicated in", r"replication sample"]),
    ("Ablation/robustness", [r"ablation", r"sensitivity analysis", r"robustness", r"permutation test", r"bootstrap"]),
    ("Statistical association/group comparison", [r"group comparison", r"correlation", r"association", r"regression analysis", r"statistical modelling", r"statistical modeling"]),
    ("Protocol/planned only", [r"protocol", r"planned", r"results not yet"]),
])

COUNTRY_PATTERNS: OrderedDict[str, Sequence[str]] = OrderedDict([
    ("Australia", [r"\baustralia\b"]), ("Austria", [r"\baustria\b"]),
    ("Belgium", [r"\bbelgium\b"]), ("Brazil", [r"\bbrazil\b"]),
    ("Canada", [r"\bcanada\b"]), ("Chile", [r"\bchile\b"]),
    ("China", [r"\bchina\b", r"chinese academy", r"beijing", r"shanghai", r"lanzhou", r"chengdu", r"wuhan", r"xi'an"]),
    ("Czech Republic", [r"czech"]), ("Denmark", [r"\bdenmark\b", r"danish"]),
    ("Finland", [r"\bfinland\b"]), ("France", [r"\bfrance\b"]),
    ("Germany", [r"\bgermany\b"]), ("Greece", [r"\bgreece\b"]),
    ("Hungary", [r"\bhungary\b"]), ("India", [r"\bindia\b"]),
    ("Iran", [r"\biran\b"]), ("Ireland", [r"\bireland\b"]),
    ("Israel", [r"\bisrael\b"]), ("Italy", [r"\bitaly\b"]),
    ("Japan", [r"\bjapan\b"]), ("Malaysia", [r"\bmalaysia\b"]),
    ("Mexico", [r"\bmexico\b"]), ("Netherlands", [r"\bnetherlands\b", r"\bdutch\b"]),
    ("New Zealand", [r"new zealand"]), ("Norway", [r"\bnorway\b"]),
    ("Pakistan", [r"\bpakistan\b"]), ("Poland", [r"\bpoland\b"]),
    ("Portugal", [r"\bportugal\b"]), ("Romania", [r"\bromania\b"]),
    ("Russia", [r"\brussia\b"]), ("Saudi Arabia", [r"saudi arabia"]),
    ("Singapore", [r"\bsingapore\b"]), ("South Korea", [r"south korea", r"\bkorea\b", r"seoul"]),
    ("Spain", [r"\bspain\b"]), ("Sweden", [r"\bsweden\b"]),
    ("Switzerland", [r"\bswitzerland\b", r"zurich"]), ("Taiwan", [r"\btaiwan\b"]),
    ("Thailand", [r"\bthailand\b", r"\bthai\b"]), ("Tunisia", [r"\btunisia\b"]),
    ("Turkey", [r"\bturkey\b", r"turkish"]),
    ("United Kingdom", [r"united kingdom", r"\buk\b", r"england", r"scotland", r"wales"]),
    ("United States", [r"united states", r"\busa\b", r"\bu\.s\.\b", r"\bus-based\b", r"california", r"new york", r"michigan", r"louisiana"]),
    ("Vietnam", [r"\bvietnam\b"]),
])

LANGUAGE_PATTERNS: OrderedDict[str, Sequence[str]] = OrderedDict([
    ("English", [r"\benglish\b"]), ("Mandarin/Chinese", [r"mandarin", r"\bchinese\b"]),
    ("German", [r"\bgerman\b"]), ("Japanese", [r"\bjapanese\b"]),
    ("Korean", [r"\bkorean\b"]), ("Danish", [r"\bdanish\b"]),
    ("Spanish", [r"\bspanish\b"]), ("French", [r"\bfrench\b"]),
    ("Italian", [r"\bitalian\b"]), ("Portuguese", [r"portuguese"]),
    ("Turkish", [r"\bturkish\b"]), ("Hebrew", [r"\bhebrew\b"]),
    ("Russian", [r"\brussian\b"]), ("Arabic", [r"\barabic\b"]),
    ("Dutch", [r"\bdutch\b"]), ("Swedish", [r"\bswedish\b"]),
    ("Norwegian", [r"\bnorwegian\b"]), ("Finnish", [r"\bfinnish\b"]),
    ("Polish", [r"\bpolish\b"]), ("Persian/Farsi", [r"persian", r"farsi"]),
    ("Bengali", [r"bengali"]), ("Thai", [r"\bthai\b"]),
    ("Hungarian", [r"\bhungarian\b"]), ("Greek", [r"\bgreek\b"]),
])

QUALITY_LABELS = OrderedDict([
    ("q_scope_clarity", "Scope clarity"),
    ("q_installation", "Installation"),
    ("q_documentation", "Documentation"),
    ("q_feature_transparency", "Feature transparency"),
    ("q_parameter_transparency", "Parameter transparency"),
    ("q_reproducible_example", "Reproducible example"),
    ("q_versioning", "Versioning"),
    ("q_testing", "Testing/CI"),
    ("q_maintenance", "Maintenance"),
    ("q_license", "Licence"),
    ("q_publication_linkage", "Publication linkage"),
    ("q_clinical_relevance", "Clinical relevance"),
])


# -----------------------------------------------------------------------------
# Utility functions
# -----------------------------------------------------------------------------


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def norm(value: Any) -> str:
    return clean_text(value).lower()


def is_missing(value: Any) -> bool:
    return norm(value) in MISSING_MARKERS


def is_not_reported(value: Any) -> bool:
    text = norm(value)
    if text in MISSING_MARKERS:
        return True
    return any(re.search(p, text, flags=re.I) for p in NOT_REPORTED_PATTERNS)


def reported_flag(value: Any) -> bool:
    return not is_not_reported(value)


def yes_no_other(value: Any) -> str:
    text = norm(value)
    if text in MISSING_MARKERS:
        return "Missing"
    if text == "yes" or text.startswith("yes;") or text.startswith("yes "):
        return "Yes"
    if text == "no" or text.startswith("no;") or text.startswith("no "):
        return "No"
    if text.startswith("other"):
        return "Other/unclear"
    return "Other/unclear"


def availability_status(value: Any) -> str:
    """Classify open code/data without treating 'on request' as public."""
    text = norm(value)
    if text in MISSING_MARKERS:
        return "Missing"
    if text == "yes" or text.startswith("yes;"):
        return "Publicly available"
    if "on request" in text or "upon request" in text or "reasonable request" in text or "qualified researcher" in text:
        return "Available on request"
    if text == "no" or text.startswith("no;"):
        return "Not publicly available"
    if "partial" in text or "package link" in text or "unclear" in text or text.startswith("other"):
        return "Partial/unclear"
    return "Partial/unclear"


def defaults_status(value: Any) -> str:
    text = norm(value)
    if text in MISSING_MARKERS:
        return "Missing"
    if text == "yes" or text.startswith("yes;"):
        return "Yes"
    if text == "no" or text.startswith("no;"):
        return "No"
    return "Partial/unclear"


def match_categories(text: str, mapping: OrderedDict[str, Sequence[str]]) -> list[str]:
    lowered = text.lower()
    found: list[str] = []
    for category, patterns in mapping.items():
        if any(re.search(p, lowered, flags=re.I) for p in patterns):
            found.append(category)
    return found


def match_dict_categories(text: str, mapping: OrderedDict[str, dict[str, Any]]) -> list[tuple[str, dict[str, Any]]]:
    lowered = text.lower()
    found: list[tuple[str, dict[str, Any]]] = []
    for category, meta in mapping.items():
        if any(re.search(p, lowered, flags=re.I) for p in meta["patterns"]):
            found.append((category, meta))
    return found


def parse_first_number(value: Any) -> float | None:
    text = clean_text(value).replace(",", "")
    if not text or is_not_reported(text):
        return None
    matches = re.findall(r"(?<!\d)(\d+(?:\.\d+)?)(?!\d)", text)
    if not matches:
        return None
    try:
        return float(matches[0])
    except ValueError:
        return None


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_int_year(value: Any) -> int | None:
    text = clean_text(value)
    if re.fullmatch(r"(?:19|20)\d{2}", text):
        return int(text)
    try:
        number = int(float(text))
        if 1900 <= number <= 2100:
            return number
    except (ValueError, TypeError):
        pass
    return None


def wrap_title(text: str, width: int = 28) -> str:
    words = text.split()
    lines: list[str] = []
    line: list[str] = []
    for word in words:
        if len(" ".join(line + [word])) > width and line:
            lines.append(" ".join(line))
            line = [word]
        else:
            line.append(word)
    if line:
        lines.append(" ".join(line))
    return "\n".join(lines)


# -----------------------------------------------------------------------------
# Study cleaning and long-format extraction
# -----------------------------------------------------------------------------


def clean_studies(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    required = ["Covidence #", "Study ID", "Year", "Diagnostic Group(s)", "Software / tool / package", "Feature(s)"]
    missing_cols = [c for c in required if c not in raw.columns]
    if missing_cols:
        raise ValueError(f"Publication CSV is missing required columns: {missing_cols}")

    repairs: list[dict[str, Any]] = []
    warnings: list[str] = []
    rows: list[dict[str, Any]] = []

    for idx, row in raw.iterrows():
        cov = clean_text(row.get("Covidence #"))
        study_key = f"COV-{cov}" if cov else f"ROW-{idx + 2}"
        year = safe_int_year(row.get("Year"))
        title_primary = clean_text(row.get("Title.1"))
        title_secondary = clean_text(row.get("Title"))

        # Some Covidence exports shift Year/Title.1 for a small number of rows.
        if year is None:
            title1_year = safe_int_year(title_primary)
            raw_year_text = clean_text(row.get("Year"))
            if title1_year is not None and len(raw_year_text) > 20:
                year = title1_year
                title_primary = raw_year_text
                repairs.append({
                    "study_key": study_key,
                    "row_number": idx + 2,
                    "field": "Year/Title.1",
                    "original_year": raw_year_text,
                    "original_title_1": clean_text(row.get("Title.1")),
                    "repair": "Swapped malformed Year and Title.1 values",
                })

        title = title_primary if title_primary and safe_int_year(title_primary) is None else title_secondary
        if not title:
            title = title_secondary

        if year is None:
            warnings.append(f"No valid year for {study_key}: {clean_text(row.get('Year'))}")

        sample_size = parse_first_number(row.get("Sample Size Total"))
        rows.append({
            "study_key": study_key,
            "covidence_number": cov,
            "study_id_raw": clean_text(row.get("Study ID")),
            "covidence_record_id": clean_text(row.get("Covidence Record ID")),
            "year": year,
            "title": title,
            "journal_source": clean_text(row.get("Journal / Source")),
            "country_setting_raw": clean_text(row.get("Country / Setting")),
            "diagnostic_groups_raw": clean_text(row.get("Diagnostic Group(s)")),
            "diagnostic_assessment_raw": clean_text(row.get("Diagnostic / symptom assessment method")),
            "sample_size_total_raw": clean_text(row.get("Sample Size Total")),
            "sample_size_total_numeric": sample_size,
            "age_raw": clean_text(row.get("Age")),
            "sex_gender_raw": clean_text(row.get("Gender/Sex Distribution")),
            "languages_raw": clean_text(row.get("Language(s)")),
            "speech_task_raw": clean_text(row.get("Speech Task Type")),
            "speech_task_details_raw": clean_text(row.get("Prompt/task details")),
            "recording_environment_raw": clean_text(row.get("Recording Environment")),
            "recording_equipment_raw": clean_text(row.get("Recording Device/eqipment")),
            "sampling_rate_raw": clean_text(row.get("Sampling rate")),
            "segmentation_vad_raw": clean_text(row.get("Segmentation / VAD")),
            "frame_length_raw": clean_text(row.get("Frame length")),
            "window_hop_raw": clean_text(row.get("window")),
            "noise_reduction_raw": clean_text(row.get("Noise Reduction")),
            "software_raw": clean_text(row.get("Software / tool / package")),
            "programming_language_raw": clean_text(row.get("Programming Language")),
            "version_reported_raw": clean_text(row.get("Version reported?")),
            "version_commit_release_raw": clean_text(row.get("Version / commit / release")),
            "features_raw": clean_text(row.get("Feature(s)")),
            "parameter_settings_raw": clean_text(row.get("Parameter Settings")),
            "defaults_assumed_raw": clean_text(row.get("Default Parameters assumed?")),
            "validation_conducted_raw": clean_text(row.get("Validation conducted?")),
            "validation_method_raw": clean_text(row.get("Validation Method")),
            "validation_details_raw": clean_text(row.get("Validation Details")),
            "open_code_raw": clean_text(row.get("Open code available?")),
            "open_data_raw": clean_text(row.get("Open data available?")),
            "repo_link_raw": clean_text(row.get("Repo Link")),
            "essential_details_missing_raw": clean_text(row.get("Essential Details Missing?")),
            "author_contact_needed_raw": clean_text(row.get("Author contact needed?")),
            "notes_raw": clean_text(row.get("Notes.1")),
            # Derived statuses
            "open_code_status": availability_status(row.get("Open code available?")),
            "open_data_status": availability_status(row.get("Open data available?")),
            "version_reported_status": yes_no_other(row.get("Version reported?")),
            "defaults_assumed_status": defaults_status(row.get("Default Parameters assumed?")),
            "validation_conducted_status": yes_no_other(row.get("Validation conducted?")),
            "essential_details_missing_status": yes_no_other(row.get("Essential Details Missing?")),
            "recording_equipment_reported": reported_flag(row.get("Recording Device/eqipment")),
            "sampling_rate_reported": reported_flag(row.get("Sampling rate")),
            "segmentation_vad_reported": reported_flag(row.get("Segmentation / VAD")),
            "frame_length_reported": reported_flag(row.get("Frame length")),
            "window_hop_reported": reported_flag(row.get("window")),
            "noise_reduction_reported": reported_flag(row.get("Noise Reduction")),
            "parameter_settings_reported": reported_flag(row.get("Parameter Settings")),
            "age_reported": reported_flag(row.get("Age")),
            "sex_gender_reported": reported_flag(row.get("Gender/Sex Distribution")),
            "diagnostic_assessment_reported": reported_flag(row.get("Diagnostic / symptom assessment method")),
        })

    studies = pd.DataFrame(rows)
    if studies["study_key"].duplicated().any():
        duplicates = studies.loc[studies["study_key"].duplicated(keep=False), "study_key"].tolist()
        raise ValueError(f"Duplicate study keys after cleaning: {duplicates[:10]}")

    return studies, pd.DataFrame(repairs), warnings


def make_long_matches(
    studies: pd.DataFrame,
    text_columns: Sequence[str],
    mapping: OrderedDict[str, Sequence[str]],
    category_column: str,
    raw_column_name: str,
    unmatched_label: str | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []
    for _, row in studies.iterrows():
        text = " ; ".join(clean_text(row.get(c)) for c in text_columns if clean_text(row.get(c)))
        categories = match_categories(text, mapping)
        if not categories and text and not is_not_reported(text):
            unmatched.append({"study_key": row["study_key"], "title": row["title"], raw_column_name: text})
            if unmatched_label:
                categories = [unmatched_label]
        for category in categories:
            rows.append({
                "study_key": row["study_key"],
                "year": row["year"],
                "title": row["title"],
                category_column: category,
                raw_column_name: text,
            })
    return pd.DataFrame(rows), pd.DataFrame(unmatched)


def make_software_long(studies: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []
    for _, row in studies.iterrows():
        text = clean_text(row["software_raw"])
        matches = match_dict_categories(text, SOFTWARE_PATTERNS)
        if not matches and text and not is_not_reported(text):
            unmatched.append({"study_key": row["study_key"], "title": row["title"], "software_raw": text})
            matches = [("Other specified", {"type": "Other/uncategorised"})]
        elif not matches and is_not_reported(text):
            matches = [("Not reported", {"type": "Missing"})]
        for software, meta in matches:
            rows.append({
                "study_key": row["study_key"], "year": row["year"], "title": row["title"],
                "software": software, "software_type": meta["type"], "software_raw": text,
            })
    return pd.DataFrame(rows).drop_duplicates(["study_key", "software"]), pd.DataFrame(unmatched)


def make_feature_long(studies: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    rows: list[dict[str, Any]] = []
    set_rows: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []
    for _, row in studies.iterrows():
        text = clean_text(row["features_raw"])
        matches = match_dict_categories(text, FEATURE_PATTERNS)
        if not matches and text and not is_not_reported(text):
            unmatched.append({"study_key": row["study_key"], "title": row["title"], "features_raw": text})
            matches = [("Other/uncategorised", {"scope": "Uncategorised"})]
        for feature, meta in matches:
            rows.append({
                "study_key": row["study_key"], "year": row["year"], "title": row["title"],
                "feature_family": feature, "feature_scope": meta["scope"], "features_raw": text,
            })
        for feature_set in match_categories(text, FEATURE_SET_PATTERNS):
            set_rows.append({
                "study_key": row["study_key"], "year": row["year"], "title": row["title"],
                "feature_set": feature_set, "features_raw": text,
            })
    return (
        pd.DataFrame(rows).drop_duplicates(["study_key", "feature_family"]),
        pd.DataFrame(set_rows).drop_duplicates(["study_key", "feature_set"]),
        pd.DataFrame(unmatched),
    )


def make_task_long(studies: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []
    for _, row in studies.iterrows():
        text = f"{clean_text(row['speech_task_raw'])}; {clean_text(row['speech_task_details_raw'])}"
        categories = match_categories(text, TASK_PATTERNS)
        if not categories:
            unmatched.append({"study_key": row["study_key"], "title": row["title"], "task_raw": text})
            categories = ["Other/unclear"]
        for category in categories:
            rows.append({
                "study_key": row["study_key"], "year": row["year"], "title": row["title"],
                "task_category": category, "task_raw": text,
            })
    return pd.DataFrame(rows).drop_duplicates(["study_key", "task_category"]), pd.DataFrame(unmatched)


def make_validation_long(studies: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for _, row in studies.iterrows():
        text = f"{clean_text(row['validation_method_raw'])}; {clean_text(row['validation_details_raw'])}; {clean_text(row['validation_conducted_raw'])}"
        categories = match_categories(text, VALIDATION_PATTERNS)
        if not categories:
            if row["validation_conducted_status"] == "No":
                categories = ["No validation reported"]
            else:
                categories = ["Other/unclear"]
        for category in categories:
            rows.append({
                "study_key": row["study_key"], "year": row["year"], "title": row["title"],
                "validation_category": category, "validation_raw": text,
            })
    return pd.DataFrame(rows).drop_duplicates(["study_key", "validation_category"])


def make_feature_sets(studies: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for _, row in studies.iterrows():
        text = clean_text(row["features_raw"])
        for feature_set in match_categories(text, FEATURE_SET_PATTERNS):
            rows.append({"study_key": row["study_key"], "year": row["year"], "title": row["title"], "feature_set": feature_set, "features_raw": text})
    return pd.DataFrame(rows).drop_duplicates(["study_key", "feature_set"])


# -----------------------------------------------------------------------------
# Repository workbook processing
# -----------------------------------------------------------------------------


def read_optional_overrides(path: Path | None, key_column: str) -> pd.DataFrame:
    if path is None:
        return pd.DataFrame()
    if not path.exists():
        raise FileNotFoundError(path)
    df = pd.read_csv(path)
    if key_column not in df.columns:
        raise ValueError(f"Override file {path} must contain {key_column!r}")
    return df


def load_repositories(
    workbook_path: Path,
    repo_overrides: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, list[str]]:
    warnings: list[str] = []
    master = pd.read_excel(workbook_path, sheet_name="Repository_Master")
    quality = pd.read_excel(workbook_path, sheet_name="Quality_Appraisal")
    links = pd.read_excel(workbook_path, sheet_name="Study_Repo_Links")

    for frame in (master, quality, links):
        if "repo_key" in frame.columns:
            frame["repo_key"] = frame["repo_key"].astype(str).str.strip().str.lower()

    consensus = master.get("consensus_decision", pd.Series(index=master.index, dtype=object)).fillna("").astype(str).str.strip()
    manual = master["manual_screening_decision"].fillna("").astype(str).str.strip()
    master["final_decision"] = consensus.where(consensus.ne(""), manual)
    master["analysis_override_note"] = ""

    if not repo_overrides.empty:
        overrides = repo_overrides.copy()
        overrides["repo_key"] = overrides["repo_key"].astype(str).str.strip().str.lower()
        override_map = overrides.set_index("repo_key")
        for idx, row in master.iterrows():
            key = row["repo_key"]
            if key not in override_map.index:
                continue
            ov = override_map.loc[key]
            if isinstance(ov, pd.DataFrame):
                ov = ov.iloc[-1]
            decision = clean_text(ov.get("final_decision"))
            if decision:
                master.at[idx, "final_decision"] = decision
            note = clean_text(ov.get("notes"))
            master.at[idx, "analysis_override_note"] = note

    included = master.loc[master["final_decision"].eq("Include")].copy()
    if included.empty:
        raise ValueError("No included repositories after final decisions/overrides")

    quality = quality.merge(
        included[["repo_key", "final_decision", "analysis_override_note"]],
        on="repo_key", how="inner", suffixes=("", "_master")
    )

    # One row per included repository. Prefer the most complete row if duplicates exist.
    quality["_nonmissing_scores"] = quality[list(QUALITY_LABELS)].notna().sum(axis=1)
    quality = quality.sort_values(["repo_key", "_nonmissing_scores"], ascending=[True, False]).drop_duplicates("repo_key")
    quality = quality.drop(columns=["_nonmissing_scores"])

    missing_quality = sorted(set(included["repo_key"]) - set(quality["repo_key"]))
    if missing_quality:
        warnings.append(f"Included repositories without a Quality_Appraisal row: {len(missing_quality)}")

    links = links.merge(
        master[["repo_key", "final_decision", "repository_class"]],
        on="repo_key", how="left", suffixes=("", "_master")
    )
    return included, quality, links, warnings


# -----------------------------------------------------------------------------
# Summary tables
# -----------------------------------------------------------------------------


def count_studies(long_df: pd.DataFrame, category: str, denominator: int) -> pd.DataFrame:
    if long_df.empty:
        return pd.DataFrame(columns=[category, "studies", "percent"])
    result = (
        long_df.drop_duplicates(["study_key", category])
        .groupby(category, dropna=False)["study_key"].nunique()
        .sort_values(ascending=False)
        .rename("studies")
        .reset_index()
    )
    result["percent"] = result["studies"] / denominator
    return result


def cross_matrix(a: pd.DataFrame, a_col: str, b: pd.DataFrame, b_col: str) -> pd.DataFrame:
    merged = a[["study_key", a_col]].drop_duplicates().merge(
        b[["study_key", b_col]].drop_duplicates(), on="study_key", how="inner"
    )
    if merged.empty:
        return pd.DataFrame()
    matrix = pd.crosstab(merged[a_col], merged[b_col])
    matrix = matrix.loc[matrix.sum(axis=1).sort_values(ascending=False).index]
    matrix = matrix[matrix.sum(axis=0).sort_values(ascending=False).index]
    return matrix


def reporting_summary(studies: pd.DataFrame) -> pd.DataFrame:
    denominator = len(studies)
    fields = OrderedDict([
        ("Diagnostic/symptom assessment", "diagnostic_assessment_reported"),
        ("Age", "age_reported"),
        ("Sex/gender", "sex_gender_reported"),
        ("Recording equipment", "recording_equipment_reported"),
        ("Sampling rate", "sampling_rate_reported"),
        ("Segmentation/VAD", "segmentation_vad_reported"),
        ("Frame length", "frame_length_reported"),
        ("Window/hop", "window_hop_reported"),
        ("Noise reduction", "noise_reduction_reported"),
        ("Parameter settings", "parameter_settings_reported"),
    ])
    rows = []
    for label, col in fields.items():
        n = int(studies[col].fillna(False).astype(bool).sum())
        rows.append({"reporting_item": label, "reported_n": n, "denominator": denominator, "percent": n / denominator})
    version_yes = int(studies["version_reported_status"].eq("Yes").sum())
    rows.append({"reporting_item": "Software version explicitly reported", "reported_n": version_yes, "denominator": denominator, "percent": version_yes / denominator})
    defaults_yes = int(studies["defaults_assumed_status"].eq("Yes").sum())
    rows.append({"reporting_item": "Defaults assumed/inferred", "reported_n": defaults_yes, "denominator": denominator, "percent": defaults_yes / denominator})
    essential_missing = int(studies["essential_details_missing_status"].eq("Yes").sum())
    rows.append({"reporting_item": "Essential methodological details missing", "reported_n": essential_missing, "denominator": denominator, "percent": essential_missing / denominator})
    return pd.DataFrame(rows).sort_values("percent", ascending=False).reset_index(drop=True)


def open_science_summary(studies: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for field, label in [("open_code_status", "Code"), ("open_data_status", "Data")]:
        counts = studies[field].value_counts(dropna=False)
        for status, count in counts.items():
            rows.append({"resource": label, "status": status, "studies": int(count), "percent": count / len(studies)})
    return pd.DataFrame(rows)


def github_quality_summary(quality: pd.DataFrame) -> pd.DataFrame:
    accessible = quality.loc[quality["repo_access_status"].astype(str).str.lower().eq("accessible")].copy()
    rows: list[dict[str, Any]] = []
    for col, label in QUALITY_LABELS.items():
        scores = pd.to_numeric(accessible[col], errors="coerce")
        denominator = int(scores.notna().sum())
        for score in [0, 1, 2]:
            n = int(scores.eq(score).sum())
            rows.append({
                "domain": label, "score": score, "repositories": n,
                "denominator": denominator, "percent": (n / denominator if denominator else np.nan),
            })
        rows.append({
            "domain": label, "score": "Missing", "repositories": int(scores.isna().sum()),
            "denominator": len(accessible), "percent": scores.isna().sum() / len(accessible) if len(accessible) else np.nan,
        })
    return pd.DataFrame(rows)


# -----------------------------------------------------------------------------
# Figures
# -----------------------------------------------------------------------------


def save_barh(df: pd.DataFrame, label_col: str, value_col: str, title: str, path: Path, top_n: int | None = None, percent: bool = False) -> None:
    data = df.copy()
    if top_n:
        data = data.head(top_n)
    data = data.sort_values(value_col, ascending=True)
    fig_height = max(4.5, 0.42 * len(data) + 1.5)
    fig, ax = plt.subplots(figsize=(10, fig_height))
    ax.barh(data[label_col].astype(str), data[value_col])
    ax.set_title(title)
    ax.set_xlabel("Percentage of studies" if percent else "Number of studies")
    if percent:
        ax.xaxis.set_major_formatter(lambda x, pos: f"{x:.0%}")
    for i, value in enumerate(data[value_col]):
        label = f"{value:.1%}" if percent else f"{int(value)}"
        ax.text(value, i, f" {label}", va="center")
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_year_chart(year_counts: pd.DataFrame, path: Path) -> None:
    data = year_counts.sort_values("year")
    fig, ax = plt.subplots(figsize=(12, 5.5))
    ax.bar(data["year"].astype(int).astype(str), data["studies"])
    ax.set_title("Included publications by year")
    ax.set_xlabel("Publication year")
    ax.set_ylabel("Number of studies")
    ax.tick_params(axis="x", rotation=60)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_heatmap(matrix: pd.DataFrame, title: str, path: Path, max_columns: int = 12) -> None:
    if matrix.empty:
        return
    matrix = matrix.iloc[:, :max_columns]
    fig_w = max(8, 0.8 * matrix.shape[1] + 4)
    fig_h = max(5, 0.55 * matrix.shape[0] + 2)
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    image = ax.imshow(matrix.values, aspect="auto")
    ax.set_xticks(range(matrix.shape[1]), labels=[wrap_title(str(x), 18) for x in matrix.columns], rotation=45, ha="right")
    ax.set_yticks(range(matrix.shape[0]), labels=[wrap_title(str(x), 28) for x in matrix.index])
    ax.set_title(title)
    for i in range(matrix.shape[0]):
        for j in range(matrix.shape[1]):
            value = matrix.iat[i, j]
            if value:
                ax.text(j, i, str(int(value)), ha="center", va="center")
    fig.colorbar(image, ax=ax, label="Number of studies")
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_quality_stacked(summary: pd.DataFrame, path: Path) -> None:
    data = summary.loc[summary["score"].isin([0, 1, 2])].copy()
    pivot = data.pivot(index="domain", columns="score", values="percent").fillna(0)
    order = pivot[2].sort_values(ascending=True).index
    pivot = pivot.loc[order]
    fig, ax = plt.subplots(figsize=(11, max(5, 0.5 * len(pivot) + 2)))
    left = np.zeros(len(pivot))
    for score in [0, 1, 2]:
        values = pivot.get(score, pd.Series(0, index=pivot.index)).values
        ax.barh(pivot.index, values, left=left, label={0: "Absent", 1: "Partial", 2: "Clearly present"}[score])
        left += values
    ax.set_title("GitHub repository quality domains")
    ax.set_xlabel("Percentage of accessible included repositories")
    ax.xaxis.set_major_formatter(lambda x, pos: f"{x:.0%}")
    ax.legend(loc="lower right")
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


# -----------------------------------------------------------------------------
# Excel output
# -----------------------------------------------------------------------------


def style_worksheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="17365D")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col_idx in range(1, ws.max_column + 1):
        header = clean_text(ws.cell(1, col_idx).value)
        values = [clean_text(ws.cell(r, col_idx).value) for r in range(1, min(ws.max_row, 200) + 1)]
        width = min(max(12, len(header) + 2, max((len(v) for v in values), default=0) + 2), 55)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def add_df_sheet(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(name[:31])
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([None if (isinstance(v, float) and math.isnan(v)) else v for v in row])
    style_worksheet(ws)
    for col_idx, col_name in enumerate(df.columns, start=1):
        if col_name == "percent" or str(col_name).endswith("_percent"):
            for r in range(2, ws.max_row + 1):
                ws.cell(r, col_idx).number_format = "0.0%"


def write_summary_workbook(path: Path, tables: dict[str, pd.DataFrame], headline: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Systematic review analysis summary"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws.merge_cells("A1:D1")
    ws.append([])
    ws.append(["Metric", "Value"])
    for key, value in headline.items():
        ws.append([key, value])
    style_worksheet(ws)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 25

    for name, df in tables.items():
        add_df_sheet(wb, name, df)
    wb.save(path)


# -----------------------------------------------------------------------------
# Main pipeline
# -----------------------------------------------------------------------------


def apply_study_overrides(studies: pd.DataFrame, overrides: pd.DataFrame) -> pd.DataFrame:
    out = studies.copy()
    out["primary_include"] = True
    out["sensitivity_group"] = "Primary"
    out["study_override_notes"] = ""
    if overrides.empty:
        return out
    ov = overrides.copy()
    ov["covidence_number"] = ov["covidence_number"].astype(str).str.strip()
    ov = ov.set_index("covidence_number")
    for idx, row in out.iterrows():
        key = str(row["covidence_number"]).strip()
        if key not in ov.index:
            continue
        x = ov.loc[key]
        if isinstance(x, pd.DataFrame):
            x = x.iloc[-1]
        include = clean_text(x.get("primary_include")).lower()
        if include in {"false", "0", "no", "exclude"}:
            out.at[idx, "primary_include"] = False
        elif include in {"true", "1", "yes", "include"}:
            out.at[idx, "primary_include"] = True
        group = clean_text(x.get("sensitivity_group"))
        if group:
            out.at[idx, "sensitivity_group"] = group
        out.at[idx, "study_override_notes"] = clean_text(x.get("notes"))
    return out


def run(args: argparse.Namespace) -> dict[str, Any]:
    studies_path = Path(args.studies).resolve()
    github_path = Path(args.github_workbook).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    long_dir = output_dir / "long_format"
    tables_dir = output_dir / "tables"
    figures_dir = output_dir / "figures"
    qa_dir = output_dir / "qa"
    for d in [long_dir, tables_dir, figures_dir, qa_dir]:
        d.mkdir(exist_ok=True)

    repo_overrides = read_optional_overrides(Path(args.repo_overrides).resolve() if args.repo_overrides else None, "repo_key")
    study_overrides = read_optional_overrides(Path(args.study_overrides).resolve() if args.study_overrides else None, "covidence_number")

    raw = pd.read_csv(studies_path)
    studies, repairs, warnings = clean_studies(raw)
    studies = apply_study_overrides(studies, study_overrides)
    primary_studies = studies.loc[studies["primary_include"]].copy()

    diagnoses, unmatched_diagnoses = make_long_matches(
        primary_studies, ["diagnostic_groups_raw"], DIAGNOSIS_PATTERNS,
        "diagnostic_group", "diagnostic_groups_raw", "Other/uncategorised",
    )
    countries, unmatched_countries = make_long_matches(
        primary_studies, ["country_setting_raw"], COUNTRY_PATTERNS,
        "country", "country_setting_raw", "Other/unclear",
    )
    languages, unmatched_languages = make_long_matches(
        primary_studies, ["languages_raw"], LANGUAGE_PATTERNS,
        "language", "languages_raw", "Other/unclear",
    )
    software, unmatched_software = make_software_long(primary_studies)
    programming_languages, unmatched_programming = make_long_matches(
        primary_studies, ["programming_language_raw"], PROGRAMMING_LANGUAGE_PATTERNS,
        "programming_language", "programming_language_raw", "Other/unclear",
    )
    features, feature_sets, unmatched_features = make_feature_long(primary_studies)
    tasks, unmatched_tasks = make_task_long(primary_studies)
    validation = make_validation_long(primary_studies)

    included_repos, repository_quality, study_repo_links, repo_warnings = load_repositories(github_path, repo_overrides)
    warnings.extend(repo_warnings)

    # Save long-format datasets.
    long_outputs = {
        "studies_clean.csv": studies,
        "study_diagnoses.csv": diagnoses,
        "study_countries.csv": countries,
        "study_languages.csv": languages,
        "study_software.csv": software,
        "study_programming_languages.csv": programming_languages,
        "study_features.csv": features,
        "study_feature_sets.csv": feature_sets,
        "study_tasks.csv": tasks,
        "study_validation.csv": validation,
        "study_repositories.csv": study_repo_links,
        "repositories_included.csv": included_repos,
        "repository_quality.csv": repository_quality,
    }
    for filename, frame in long_outputs.items():
        frame.to_csv(long_dir / filename, index=False)

    # Summary tables.
    n = len(primary_studies)
    year_counts = (
        primary_studies.dropna(subset=["year"]).groupby("year")["study_key"].nunique().rename("studies").reset_index().sort_values("year")
    )
    diagnosis_counts = count_studies(diagnoses, "diagnostic_group", n)
    country_counts = count_studies(countries, "country", n)
    language_counts = count_studies(languages, "language", n)
    software_counts = count_studies(software, "software", n)
    programming_counts = count_studies(programming_languages, "programming_language", n)
    feature_counts = count_studies(features.loc[features["feature_scope"].ne("Non-acoustic/multimodal")], "feature_family", n)
    feature_set_counts = count_studies(feature_sets, "feature_set", n)
    task_counts = count_studies(tasks, "task_category", n)
    validation_counts = count_studies(validation, "validation_category", n)
    report_counts = reporting_summary(primary_studies)
    open_science = open_science_summary(primary_studies)
    diagnosis_feature = cross_matrix(diagnoses, "diagnostic_group", features.loc[features["feature_scope"].ne("Non-acoustic/multimodal")], "feature_family")
    diagnosis_software = cross_matrix(diagnoses, "diagnostic_group", software.loc[~software["software"].isin(["Other specified", "Not reported"])], "software")

    repo_class = (
        included_repos.groupby("repository_class", dropna=False)["repo_key"].nunique().rename("repositories").reset_index().sort_values("repositories", ascending=False)
    )
    repo_class["percent"] = repo_class["repositories"] / len(included_repos)
    readme_status = (
        repository_quality.groupby("readme_status", dropna=False)["repo_key"].nunique().rename("repositories").reset_index().sort_values("repositories", ascending=False)
    )
    readme_status["percent"] = readme_status["repositories"] / len(repository_quality)
    access_status = (
        repository_quality.groupby("repo_access_status", dropna=False)["repo_key"].nunique().rename("repositories").reset_index().sort_values("repositories", ascending=False)
    )
    access_status["percent"] = access_status["repositories"] / len(repository_quality)
    quality_summary = github_quality_summary(repository_quality)

    quality_accessible = repository_quality.loc[repository_quality["repo_access_status"].astype(str).str.lower().eq("accessible")].copy()
    quality_total = pd.to_numeric(quality_accessible.get("quality_total_scored"), errors="coerce")
    quality_stats = pd.DataFrame([
        {"statistic": "Accessible repositories", "value": len(quality_accessible)},
        {"statistic": "Mean quality total", "value": quality_total.mean()},
        {"statistic": "Median quality total", "value": quality_total.median()},
        {"statistic": "Q1 quality total", "value": quality_total.quantile(0.25)},
        {"statistic": "Q3 quality total", "value": quality_total.quantile(0.75)},
        {"statistic": "Minimum quality total", "value": quality_total.min()},
        {"statistic": "Maximum quality total", "value": quality_total.max()},
    ])

    table_outputs: dict[str, pd.DataFrame] = {
        "publication_years": year_counts,
        "diagnoses": diagnosis_counts,
        "countries": country_counts,
        "languages": language_counts,
        "software": software_counts,
        "programming_languages": programming_counts,
        "feature_families": feature_counts,
        "feature_sets": feature_set_counts,
        "speech_tasks": task_counts,
        "validation": validation_counts,
        "reporting_completeness": report_counts,
        "open_science": open_science,
        "github_repository_classes": repo_class,
        "github_readme_status": readme_status,
        "github_access_status": access_status,
        "github_quality_domains": quality_summary,
        "github_quality_total": quality_stats,
    }
    for name, frame in table_outputs.items():
        frame.to_csv(tables_dir / f"{name}.csv", index=False)
    diagnosis_feature.to_csv(tables_dir / "diagnosis_feature_matrix.csv")
    diagnosis_software.to_csv(tables_dir / "diagnosis_software_matrix.csv")

    # QA/audit outputs.
    qa_outputs = {
        "repairs_log.csv": repairs,
        "unmatched_diagnoses.csv": unmatched_diagnoses,
        "unmatched_countries.csv": unmatched_countries,
        "unmatched_languages.csv": unmatched_languages,
        "unmatched_software.csv": unmatched_software,
        "unmatched_programming_languages.csv": unmatched_programming,
        "unmatched_features.csv": unmatched_features,
        "unmatched_tasks.csv": unmatched_tasks,
    }
    for filename, frame in qa_outputs.items():
        frame.to_csv(qa_dir / filename, index=False)

    # Figures.
    save_year_chart(year_counts, figures_dir / "01_publications_by_year.png")
    save_barh(diagnosis_counts, "diagnostic_group", "studies", "Studies by psychiatric diagnostic group", figures_dir / "02_diagnostic_groups.png")
    save_barh(feature_counts, "feature_family", "studies", "Acoustic feature families", figures_dir / "03_feature_families.png")
    save_barh(software_counts.loc[~software_counts["software"].isin(["Other specified", "Not reported"])], "software", "studies", "Most frequently reported software and toolkits", figures_dir / "04_software_tools.png", top_n=15)
    save_barh(report_counts.sort_values("percent", ascending=False), "reporting_item", "percent", "Methodological reporting completeness", figures_dir / "05_reporting_completeness.png", percent=True)
    save_barh(validation_counts, "validation_category", "studies", "Validation and evaluation approaches", figures_dir / "06_validation_methods.png")
    save_heatmap(diagnosis_feature, "Diagnostic group × acoustic feature family", figures_dir / "07_diagnosis_feature_heatmap.png")
    save_heatmap(diagnosis_software, "Diagnostic group × software/tool", figures_dir / "08_diagnosis_software_heatmap.png")
    save_quality_stacked(quality_summary, figures_dir / "09_github_quality_domains.png")
    save_barh(readme_status, "readme_status", "repositories", "README status among included repositories", figures_dir / "10_github_readme_status.png")

    # Summary workbook.
    headline = OrderedDict([
        ("Included publications", n),
        ("Publications from 2020 onward", int(primary_studies["year"].ge(2020).sum())),
        ("Included repositories", len(included_repos)),
        ("Accessible included repositories", len(quality_accessible)),
        ("Repositories with README retrieved", int(repository_quality["readme_status"].eq("Retrieved").sum())),
        ("Repositories with confirmed absent README", int(repository_quality["readme_status"].eq("Confirmed absent").sum())),
        ("Repository quality mean /24", round(float(quality_total.mean()), 2) if quality_total.notna().any() else ""),
        ("Repository quality median /24", round(float(quality_total.median()), 2) if quality_total.notna().any() else ""),
    ])
    workbook_tables = {
        "Publication_years": year_counts,
        "Diagnoses": diagnosis_counts,
        "Software": software_counts,
        "Feature_families": feature_counts,
        "Speech_tasks": task_counts,
        "Validation": validation_counts,
        "Reporting": report_counts,
        "Open_science": open_science,
        "GitHub_classes": repo_class,
        "GitHub_README": readme_status,
        "GitHub_quality": quality_summary,
        "GitHub_quality_total": quality_stats,
        "Dx_feature_matrix": diagnosis_feature.reset_index(),
        "Dx_software_matrix": diagnosis_software.reset_index(),
    }
    workbook_path = output_dir / "systematic_review_analysis_summary.xlsx"
    write_summary_workbook(workbook_path, workbook_tables, headline)

    manifest = {
        "run_timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "python_version": sys.version,
        "inputs": {
            "studies": {"path": str(studies_path), "sha256": file_sha256(studies_path)},
            "github_workbook": {"path": str(github_path), "sha256": file_sha256(github_path)},
            "repo_overrides": str(Path(args.repo_overrides).resolve()) if args.repo_overrides else None,
            "study_overrides": str(Path(args.study_overrides).resolve()) if args.study_overrides else None,
        },
        "counts": {
            "raw_publication_rows": len(raw),
            "primary_included_publications": n,
            "diagnosis_long_rows": len(diagnoses),
            "software_long_rows": len(software),
            "feature_long_rows": len(features),
            "task_long_rows": len(tasks),
            "included_repositories": len(included_repos),
            "accessible_included_repositories": len(quality_accessible),
        },
        "repairs": len(repairs),
        "warnings": warnings,
    }
    with (output_dir / "analysis_manifest.json").open("w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)

    summary_lines = [
        "# Systematic review analysis run",
        "",
        f"- Included publications: **{n}**",
        f"- Publications from 2020 onward: **{int(primary_studies['year'].ge(2020).sum())} ({primary_studies['year'].ge(2020).mean():.1%})**",
        f"- Included GitHub repositories: **{len(included_repos)}**",
        f"- Accessible included repositories: **{len(quality_accessible)}**",
        f"- README retrieved: **{int(repository_quality['readme_status'].eq('Retrieved').sum())}**",
        f"- Confirmed absent README: **{int(repository_quality['readme_status'].eq('Confirmed absent').sum())}**",
        f"- Mean repository quality total: **{quality_total.mean():.2f}/24**",
        f"- Median repository quality total: **{quality_total.median():.1f}/24**",
        "",
        "## Important audit notes",
        f"- Automated Year/Title repairs: {len(repairs)}",
        f"- Unmatched software descriptions requiring review: {len(unmatched_software)}",
        f"- Unmatched feature descriptions requiring review: {len(unmatched_features)}",
        f"- Unmatched task descriptions requiring review: {len(unmatched_tasks)}",
        "- Review the files in `qa/` before treating all harmonised categories as final.",
    ]
    if warnings:
        summary_lines.extend(["", "## Warnings"] + [f"- {w}" for w in warnings])
    (output_dir / "analysis_summary.md").write_text("\n".join(summary_lines), encoding="utf-8")

    return {
        "output_dir": str(output_dir),
        "workbook": str(workbook_path),
        "manifest": str(output_dir / "analysis_manifest.json"),
        "summary": str(output_dir / "analysis_summary.md"),
        "headline": headline,
        "warnings": warnings,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--studies", required=True, help="Latest Covidence publication extraction CSV")
    parser.add_argument("--github-workbook", required=True, help="Quality-reviewed GitHub workbook XLSX")
    parser.add_argument("--output-dir", default="systematic_review_analysis_outputs", help="Output directory")
    parser.add_argument("--repo-overrides", help="Optional CSV: repo_key, final_decision, notes")
    parser.add_argument("--study-overrides", help="Optional CSV: covidence_number, primary_include, sensitivity_group, notes")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    result = run(args)
    print(json.dumps(result, indent=2, default=str))


if __name__ == "__main__":
    main()
